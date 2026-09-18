# -*- coding: utf-8 -*-
"""대형 기획전 자동화 (추석 기획전 타입)

  python be.py init   <기획전폴더> [--open 2026-10-01] [--code HSH]   폴더 뼈대 · 이미지 목록 · 기본 정보(대형기획전.json)
  python be.py make   <기획전폴더>                                    기획 엑셀 → settings_bigevent.js + 게시판 붙여넣기 index.html
  python be.py check  <기획전폴더>                                    상품번호 · 날짜 · 데일리 개수 · 이미지 점검
  python be.py render <기획전폴더> [게시글주소]                        PC·모바일 화면 검사 (주소 없으면 최근 기획전 글에 끼워서)
  python be.py all    <기획전폴더>                                    check + render

  · 모양·기능 = 공통 파일(engine/big_event.css · big_event.js, CDN 2026/00/bigevent/)
  · 기획전마다 = settings_bigevent.js(CDN 기획전 폴더) + index.html(게시판에 붙여넣기)
  · 엑셀에 없는 값은 기본값(samples/추석)을 쓰고 '확인 필요'로 알린다. 빈 섹션은 통째로 뺀다.
  · 기존 파일은 덮지 않는다 — 이미 있으면 _자동생성/ 에 만든다.
"""
import io, os, re, sys, json, copy, glob, shutil, subprocess, datetime, urllib.request
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

KIT = Path(__file__).resolve().parent
BASE = KIT / 'samples' / '추석'
CDN = 'https://cdn-tgreen.bizhost.kr/phpskr/tgreen09/'
ENGINE_VER = '260918'
PROJECT = '대형기획전.json'
SECTIONS = ['coupon', 'gift', 'luckybag', 'daily01', 'daily02', 'review', 'lotto', 'ranking', 'membership']
QUICK_OF = {'#coupon_section':'coupon', '#gift_section':'gift', '#luckybag_section':'luckybag', '#daily_section':'daily01',
            '#daily_section02':'daily02', '#review_section':'review', '#lotto_section':'lotto'}

def rd(p): return io.open(p, encoding='utf-8').read()
def wr(p, s):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    io.open(p, 'w', encoding='utf-8', newline='').write(s)
def lines(v): return [l.strip() for l in str(v or '').replace('\r', '').split('\n')]
def nonempty(v): return [l for l in lines(v) if l]

# ───────────────────────── init
def cmd_init(folder, args):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    opt = dict(zip(args[::2], args[1::2]))
    pj = folder / PROJECT
    info = json.loads(rd(pj)) if pj.exists() else {}
    info.setdefault('open', opt.get('--open', ''))
    info.setdefault('code', opt.get('--code', 'HSH'))
    wr(pj, json.dumps(info, ensure_ascii=False, indent=1))
    for d in ['img/top', 'img/pd/title', 'img/quick', 'img/giftcard/id', 'img/luckybag/id', 'img/pd', '_소재']:
        (folder / d).mkdir(parents=True, exist_ok=True)
    wr(folder / '이미지_목록.md', image_list())
    print('만듦 :', folder)
    print('  대형기획전.json  오픈일 %s · 폴더코드 %s  (오픈일이 비면 make 때 엑셀 기간 시작일로)' % (info['open'] or '(엑셀에서)', info['code']))
    print('  img/…  이미지 넣을 곳 · 이미지_목록.md = 파일 이름·크기 (추석과 같은 이름으로 넣으면 그대로 들어감)')

def image_refs(S):
    """설정이 실제로 쓰는 이미지 : [(img/ 아래 경로, 쓰는 곳)]"""
    img = S['img']; refs = []
    def walk(v, where):
        if isinstance(v, dict):
            for k, x in v.items(): walk(x, where)
        elif isinstance(v, list):
            for x in v: walk(x, where)
        elif isinstance(v, str) and v.startswith(img): refs.append((v[len(img):], where))
    for key, where in [('config', '메인 비주얼'), ('gift', '상품권'), ('luckybag', '럭키백'), ('dailyVisual', '데일리 배너')]:
        walk(S.get(key), where)
    for k, t in S['titleImg'].items(): refs.append(('pd/title/' + t, '섹션 제목(%s)' % k))
    for q in (S.get('quick') or {}).get('items', []): refs.append(('quick/' + q['img'], '퀵 메뉴(%s)' % q.get('name', '')))
    seen = set()
    return [(r, w) for r, w in refs if not (r in seen or seen.add(r))]

def image_list():
    """이미지 목록 — 설정이 쓰는 경로 그대로(= CDN img/ 아래 경로), 크기는 추석 기준(samples/추석/image_sizes.json)"""
    S = json.loads(rd(BASE / 'settings.json'))
    sizes = json.loads(rd(BASE / 'image_sizes.json')) if (BASE / 'image_sizes.json').exists() else {}
    out = ['# 이미지 목록 — `img/` 아래 이 경로 · 이 이름 그대로 (CDN 에도 같은 경로로 올림)', '',
           '| 경로 (img/ 아래) | 쓰는 곳 | 크기(px, 추석 기준) |', '|---|---|---|']
    for rel, where in image_refs(S):
        out.append('| %s | %s | %s |' % (rel, where, sizes.get(Path(rel).name, '-')))
    out += ['', '- 섹션을 안 쓰면 그 섹션 이미지는 필요 없다.', '- 섹션 제목 이미지(pd/title/tt_*.png)가 없으면 글자 제목으로 나온다.',
            '- 크기가 달라지면(특히 메인 비주얼 제목 이미지) settings 의 위치(left/top/width)를 맞춘다.']
    return '\n'.join(out) + '\n'

# ───────────────────────── 엑셀 읽기
def find_excel(folder):
    for f in sorted(glob.glob(str(Path(folder) / '*.xls*'))):
        if Path(f).name.startswith('~$'): continue
        return f
    return None

def load_sheet(path):
    """기획파일 시트 → 행 목록 [{col: 값}] (열은 0부터, 날짜는 datetime)"""
    rows = []
    if path.lower().endswith('.xls'):
        try:
            import xlrd
        except ImportError:      # 옛 .xls 는 xlrd 가 필요 — 없는 PC 면 한 번 설치
            print('xlrd 설치 중… (옛 .xls 엑셀을 읽는 데 필요, 처음 한 번)')
            subprocess.call([sys.executable, '-m', 'pip', 'install', '--quiet', 'xlrd'])
            import xlrd
        book = xlrd.open_workbook(path)
        sh = next((s for s in book.sheets() if '기획' in s.name), book.sheet_by_index(0))
        for r in range(sh.nrows):
            row = {}
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.value in ('', None): continue
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate_as_datetime(v, book.datemode)
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                row[c] = v
            rows.append(row)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = next((w for w in wb.worksheets if '기획' in w.title), wb.worksheets[0])
        for r in ws.iter_rows():
            row = {}
            for cell in r:
                v = cell.value
                if v in ('', None): continue
                if isinstance(v, float) and v.is_integer(): v = int(v)
                row[cell.column - 1] = v
            rows.append(row)
    return rows

C, G, H, T, X, Y = 2, 6, 7, 19, 23, 24

def kind_of(label, row):
    lab = re.sub(r'\s+', '', label)
    if row.get(H) == '날짜': return 'daily'
    for key, words in [('main', ['메인타이틀']), ('sub', ['서브타이틀']), ('brands', ['관련브랜드']), ('link', ['참고링크']),
                       ('coupon', ['쿠폰']), ('luckybag', ['럭키백']), ('gift', ['예치금', '상품권']), ('lotto', ['로또']),
                       ('review', ['후기']), ('membership', ['가족', '가입', '멤버십', '회원'])]:
        if any(w in lab for w in words): return key
    return 'unknown'

def xl_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, (int, float)) and 30000 < v < 80000:
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date()
    return None

def parse_plan(rows):
    """기획파일 → 섹션별 값"""
    starts = [i for i, r in enumerate(rows) if str(r.get(C, '')).strip()]
    plan = {'daily': [], 'unknown': []}
    for n, i in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(rows)
        block = rows[i:end]
        label = str(block[0].get(C, ''))
        kind = kind_of(label, block[0])
        sec = {'label': label, 'labelLines': nonempty(label), 'G': str(block[0].get(G, '')), 'T': str(block[0].get(T, '')), 'rows': block, 'at': i + 1}
        if kind == 'daily': plan['daily'].append(sec)
        elif kind == 'unknown': plan['unknown'].append(sec)
        else: plan[kind] = sec
    return plan

def notice_items(t):
    return [re.sub(r'^-\s*', '', l) for l in nonempty(t) if l.startswith('-')]

def speaker_join(ls):
    return ' '.join(re.sub(r'^([^:：]{1,12})\s*[:：]\s*', r'\1 : ', l, count=1) for l in ls)

SPEAKER = re.compile(r'^[^:：\s][^:：]{0,11}\s*[:：]')

def copy_lines(g):
    """G 칸 대사 : 빈 줄 앞 덩어리. '이모:' 처럼 말하는 사람이 있으면 그 줄들만"""
    out = []
    for l in lines(g):
        if not l:
            if out: break
            continue
        out.append(l)
    if any(SPEAKER.match(l) for l in out):
        out = [l for l in out if SPEAKER.match(l)]
    return out

def daily_grid(sec):
    """날짜 표(날짜·상품번호·브랜드 세 칸씩) → {날짜: [상품번호…]}"""
    data = {}
    rows = sec['rows']
    for hi, r in enumerate(rows):
        if r.get(H) != '날짜': continue
        cols = [c for c, v in r.items() if v == '날짜']
        for c in cols:
            day = None; items = []
            for rr in rows[hi + 1:]:
                if rr.get(H) == '날짜': break
                d = xl_date(rr.get(c))
                if d and day is None: day = d
                p = rr.get(c + 1)
                if isinstance(p, int) or (isinstance(p, str) and p.strip().isdigit()):
                    items.append(str(int(p)))
            if day and items: data[day.isoformat()] = items
    return dict(sorted(data.items()))

# ───────────────────────── 설정 · 글자 만들기
def emit(v, img_old, ind=2):
    """파이썬 값 → JS 글자. 이미지 주소는 IMG + '…'"""
    pad = ' ' * ind
    if isinstance(v, dict):
        if not v: return '{}'
        return '{\n' + ',\n'.join('%s  %s: %s' % (pad, k if re.match(r'^[A-Za-z_$][\w$]*$', k) else json.dumps(k, ensure_ascii=False), emit(x, img_old, ind + 2)) for k, x in v.items()) + '\n' + pad + '}'
    if isinstance(v, list):
        if not v: return '[]'
        if all(not isinstance(x, (dict, list)) for x in v):
            return '[' + ', '.join(emit(x, img_old, ind) for x in v) + ']'
        return '[\n' + ',\n'.join('%s  %s' % (pad, emit(x, img_old, ind + 2)) for x in v) + '\n' + pad + ']'
    if isinstance(v, str) and img_old and v.startswith(img_old):
        return 'IMG + ' + json.dumps(v[len(img_old):], ensure_ascii=False)
    return json.dumps(v, ensure_ascii=False)

def build(folder):
    folder = Path(folder)
    base = json.loads(rd(BASE / 'settings.json'))
    texts = json.loads(rd(BASE / 'texts.json'))
    S = copy.deepcopy(base); TX = dict(texts)
    got, keep = [], []              # 엑셀에서 읽은 것 / 기본값 그대로(확인 필요)
    xl = find_excel(folder)
    if not xl: raise SystemExit('★ 기획 엑셀(.xls/.xlsx)이 폴더에 없음 : %s' % folder)
    plan = parse_plan(load_sheet(xl))
    info = json.loads(rd(folder / PROJECT)) if (folder / PROJECT).exists() else {}

    # 기간 · 이름
    sub = plan.get('sub', {})
    m = re.search(r'(\d{4})\.(\d{1,2})\.(\d{1,2})\s*[-~]\s*(?:(\d{4})\.)?(\d{1,2})\.(\d{1,2})', sub.get('T', '') + ' ' + sub.get('G', ''))
    if not m: raise SystemExit('★ 서브 타이틀 칸에서 기간(2026.09.14 - 2026.09.27)을 못 찾음')
    y1, m1, d1, y2, m2, d2 = m.groups(); y2 = y2 or y1
    start = datetime.date(int(y1), int(m1), int(d1)); end = datetime.date(int(y2), int(m2), int(d2))
    open_day = info.get('open') or start.isoformat()
    od = datetime.date.fromisoformat(open_day)
    code = info.get('code') or 'HSH'
    folder_url = '%s%04d/%02d/%02d/%s/' % (CDN, od.year, od.month, od.day, code)
    img_old, IMG = base['img'], folder_url + 'img/'
    S['img'] = IMG
    main = plan.get('main', {})
    mg = nonempty(main.get('G', ''))
    name_m = re.search(r'\((.+?)\s*기획전\)', main.get('G', ''))
    name = name_m.group(1) if name_m else Path(folder).name.split('_', 1)[-1].replace('기획전', '').strip()
    title = mg[0] if mg else ''
    upto = re.search(r'up\s*to\s*(\d+)\s*%', sub.get('T', '') + sub.get('G', ''), re.I)
    S['config']['eventStart'] = '%sT%02d:00:00+09:00' % (start.isoformat(), S['config'].get('openHour', 10))
    S['config']['eventEnd'] = '%sT23:59:59+09:00' % end.isoformat()
    S['config']['visual']['alt'] = '%d %s 기획전 - %s' % (start.year, name, title.rstrip('.'))
    for dev in ('pc', 'mobile'):
        for L in S['config']['visual'][dev]['layers']:
            keep.append('비주얼 제목 이미지 위치(%s) — 새 제목 이미지면 left/top/width 조정' % dev) if False else None
    TX['period'] = '%s - %02d.%02d' % (start.strftime('%Y.%m.%d'), end.month, end.day)
    got.append('기간 %s ~ %s · 이름 "%s" · 메인 타이틀 "%s"%s' % (start, end, name, title, (' · UP TO %s%%' % upto.group(1)) if upto else ''))
    keep.append('비주얼 제목 이미지(top/img_top_title01·02.png) 위치 — 추석 값 그대로. 새 제목 이미지 크기가 다르면 settings 의 visual.layers left/top/width 조정')
    S['relatedButton']['href'] = ''
    keep.append("'전 상품 보기' 버튼 — 링크 비워 둠(게시글 번호는 매번 달라서). 쓰려면 settings 의 relatedButton.href 에 주소")

    off = set()
    notices = {n['id']: n for n in S['noticeModals']}
    def set_notice(nid, items, what):
        if not items: keep.append('%s 유의사항 — 엑셀에 없어 추석 문구 그대로' % what); return
        old = notices[nid]['items']
        plain = lambda h: re.sub(r'<[^>]+>', '', h).replace('&gt;', '>').replace('&lt;', '<').replace(' ', '')
        new = [next((o for o in old if plain(o) == plain(i)), i) for i in items]      # 같은 문장이면 기본값(링크 등 포함)을 씀
        notices[nid]['items'] = new
        got.append('%s 유의사항 %d줄' % (what, len(items)))
        dropped = [re.sub(r'<[^>]+>', '', o) for o in old if o not in new]
        if dropped:
            keep.append('%s 유의사항 — 추석 페이지에만 있던 %d줄은 엑셀에 없어 뺐음(필요하면 settings 에 추가): %s' % (what, len(dropped), ' / '.join(d[:40] for d in dropped)))

    # 쿠폰
    cp = plan.get('coupon')
    if cp:
        specials = []
        for r in cp['rows']:
            rate = re.search(r'(\d+)\s*%', str(r.get(X, ''))); idx = re.search(r'idx=(\d+)', str(r.get(Y, '')))
            if rate and idx:
                specials.append({'rate': rate.group(1) + '%', 'sub': '%s 특별 쿠폰' % name, 'idx': idx.group(1), 'className': 'is_fixed', 'soldout': False})
        S['coupon']['specials'] = specials; S['coupon']['daily']['items'] = []
        TX['coupon_title'] = ' '.join(cp['labelLines'])
        cl = copy_lines(cp['G']); TX['coupon_pc'] = speaker_join(cl); TX['coupon_mo'] = speaker_join(cl[:1])
        set_notice('coupon_notice_modal', notice_items(cp['T']), '쿠폰')
        got.append('쿠폰 %d장 (%s)' % (len(specials), ', '.join('%s idx %s' % (s['rate'], s['idx']) for s in specials)))
        if not specials: off.add('coupon')
    else: off.add('coupon')

    # 상품권
    gf = plan.get('gift')
    if gf:
        items = []
        for r in gf['rows']:
            pno, nm = r.get(X), str(r.get(Y, ''))
            if not (isinstance(pno, int) and nm): continue
            amt = re.search(r'(\d+)\s*만\s*원', nm); rate = re.search(r'(\d+)\s*%', nm)
            won = int(amt.group(1)) if amt else 0
            items.append({'title': re.sub(r'\s*\(\s*\d+\s*%\s*할인\s*\)\s*$', '', nm), 'hint': '%d만원 상품권 구매하기' % won, 'href': '/shop/view.php?index_no=%d' % pno,
                          'image': IMG + 'giftcard/id/img_giftcard_thum_id%d.png' % won, 'originalPrice': won * 10000, 'discountRate': int(rate.group(1)) if rate else 0})
        S['gift'] = items
        TX['gift_title'] = ' '.join(gf['labelLines'])
        cl = copy_lines(gf['G']); TX['gift_pc'] = speaker_join(cl); TX['gift_mo'] = speaker_join(cl[:1])
        set_notice('gift_notice_modal', notice_items(gf['T']), '상품권(예치금)')
        got.append('상품권 %d종 (%s)' % (len(items), ', '.join('%s %d%%' % (i['href'].split('=')[-1], i['discountRate']) for i in items)))
        if not items: off.add('gift')
    else: off.add('gift')

    # 럭키백
    lb = plan.get('luckybag')
    if lb:
        r0 = next((r for r in lb['rows'] if isinstance(r.get(X), int)), None)
        if r0:
            pno, nm = r0[X], str(r0.get(Y, '')).strip()
            price = re.search(r'([\d,]+)\s*원에?\s*판매', lb['T'])
            L = S['luckybag']
            L['href'] = '/shop/view.php?index_no=%d' % pno; L['alt'] = nm; L['soldout'] = 'auto'
            if price: L['priceText'] = price.group(1) + '원'
            else: keep.append('럭키백 가격 — 엑셀 T칸에 "…원에 판매" 가 없어 추석 값 그대로')
            L['stage']['base']['alt'] = nm
            L['buttons'] = [{'label': '%s 구매하기' % nm, 'href': L['href'], 'hint': '%s 구매하기' % nm, 'soldout': False}]
            TX['luckybag_title'] = ' '.join(lb['labelLines'])
            cl = copy_lines(lb['G']); TX['luckybag_pc'] = speaker_join(cl); TX['luckybag_mo'] = speaker_join(cl[:1])
            got.append('럭키백 %d · %s · %s' % (pno, nm, L['priceText']))
            keep.append('럭키백 그림 위 글자 이미지(luckybag/id/page_luckybag_thum_id01·02.png)와 위치 — 추석 값 그대로')
        else: off.add('luckybag')
    else: off.add('luckybag')

    # 데일리 (순서대로 daily_01 · daily_02)
    for n, key in enumerate(['daily_01', 'daily_02']):
        tk = key.replace('_', '')
        if n < len(plan['daily']):
            sec = plan['daily'][n]
            data = daily_grid(sec)
            S['dailyData'][key] = data
            ll = sec['labelLines']
            TX[tk + '_title'] = ll[0]
            q = [l.lstrip("'‘’\"") for l in ll[1:]]
            if q: TX[tk + '_copy'] = ' '.join(q)
            else: keep.append('%s 문구 — 엑셀에 없어 추석 문구 그대로' % ll[0])
            S['dailyVisual'][key]['alt'] = ll[0]
            got.append('%s "%s" %d일 × %s개' % (key, ll[0], len(data), '/'.join(sorted({str(len(v)) for v in data.values()})) or '0'))
            if not data: off.add(tk)
        else:
            off.add(tk); S['dailyData'][key] = {}

    # 후기
    rv = plan.get('review')
    if rv:
        TX['review_title'] = ' '.join(rv['labelLines'])
        cl = copy_lines(rv['G']); TX['review_pc'] = speaker_join(cl); TX['review_mo'] = speaker_join(cl[:1])
        amt = re.search(r'기존\s*([\d,]+)\s*원\s*[→>-]+\s*([\d,]+)\s*원', rv['G'])
        if amt:
            TX['review_before'] = amt.group(1) + '원'; TX['review_amount'] = amt.group(2)
            TX['review_past_pc'] = re.sub(r'<em>1건당 [\d,]+원</em>', '<em>1건당 %s원</em>' % amt.group(2), TX['review_past_pc'])
            TX['review_past_mo'] = re.sub(r'<em>1건당 [\d,]+원</em>', '<em>1건당 %s원</em>' % amt.group(2), TX['review_past_mo'])
            got.append('후기 %s원 → %s원' % (amt.group(1), amt.group(2)))
        else: keep.append('후기 금액 — 엑셀에 "기존 1,000원 → 10,000원" 형식이 없어 추석 값 그대로')
        pay = end + datetime.timedelta(days=1)
        TX['review_write_pc'] = '%s ~ %s' % (start.strftime('%Y.%m.%d'), end.strftime('%Y.%m.%d'))
        TX['review_write_mo'] = '%02d.%02d ~ %02d.%02d' % (start.month, start.day, end.month, end.day)
        TX['review_pay_pc'] = '%s 일괄 지급' % pay.strftime('%Y.%m.%d'); TX['review_pay_mo'] = '%02d.%02d 일괄 지급' % (pay.month, pay.day)
        keep.append('후기 작은 글("%s") — 이벤트 이름이 들어가면 확인' % re.sub('<[^>]+>', '', TX['review_sub']))
        set_notice('review_notice_modal', notice_items(rv['T']), '후기(적립금)')
    else: off.add('review')

    # 로또
    lt = plan.get('lotto')
    if lt:
        prize = re.search(r'예치금\s*([\d,]+)\s*만\s*원', lt['G'])
        if prize:
            man = int(prize.group(1).replace(',', ''))
            S['prizeModal']['reward'] = '{:,}'.format(man * 10000)
            S['prizeModal']['notice'] = '예치금 %d만원 · 최종 당첨자 1명 지급' % man
            TX['band_title'] = TX['lotto_title'] = '로또 %d만원' % man
            S['quick']['band']['title'] = TX['band_title']
            got.append('로또 경품 예치금 %d만원' % man)
        else: keep.append('로또 경품 금액 — 엑셀에 "예치금 ○○만원" 이 없어 추석 값 그대로')
        keep.append('로또 소개 글(PC/모바일) — 추석 문구 그대로("%s")' % TX['lotto_pc'][:30])
        S['winnerModal']['released'] = False
        set_notice('lotto_notice_modal', notice_items(lt['T']), '로또')
    else: off.add('lotto')

    # 멤버십 · 랭킹 : 매번 같은 내용 → 기본값
    if 'membership' not in plan: off.add('membership')
    keep.append('회원 혜택 4장 · 랭킹 문구 — 추석(4XR 공통 혜택) 그대로')

    # 퀵 메뉴 · 스티키 : 빠진 섹션 정리
    S['quick']['items'] = [q for q in S['quick']['items'] if QUICK_OF.get(q['href']) not in off]
    for q in S['quick']['items']:
        sec = QUICK_OF.get(q['href'])
        if sec and TX.get(sec + '_title'):
            q['alt'] = q['name'] = TX[sec + '_title']
    keep.append('퀵 메뉴 짧은 이름(nameMo)·설명(desc)·아이콘 — 추석 문구 그대로, 필요하면 settings 의 quick.items')
    if 'lotto' in off: S['quick']['band'] = None      # 퀵 로또 띠도 없앰 (화면 틀에서도 빠짐)
    return S, TX, off, got, keep, {'xl': xl, 'folder_url': folder_url, 'img_old': img_old, 'name': name, 'start': start, 'end': end}

def settings_js(S, meta):
    order = ['img', 'config', 'relatedButton', 'coupon', 'gift', 'luckybag', 'noticeModals', 'prizeModal', 'winnerModal', 'lottoUi', 'comment',
             'membership', 'quick', 'titleImg', 'dailyVisual', 'dailyData']
    body = ',\n'.join('    %s: %s' % (k, emit(S[k], meta['img_old'], 4) if k != 'img' else 'IMG') for k in order)
    return ('/* 4XR 대형 기획전 설정 — %s (%s ~ %s). be.py make 로 만듦. 이 파일만 기획전마다 바뀐다 */\n'
            '(function(){\n  const IMG = %s;\n  window.BF_EVENT = {\n%s\n  };\n})();\n') % (meta['name'], meta['start'], meta['end'], json.dumps(meta['folder_url'] + 'img/'), body)

def index_html(TX, off, meta, ver):
    tpl = rd(KIT / 'templates' / 'markup.html')
    for k in sorted(off) + (['related'] if False else []):
        tpl = re.sub(r'<!--@%s-->.*?<!--@/%s-->\n' % (k, k), '', tpl, flags=re.S)
    tpl = re.sub(r'(?m)^<!--@/?[a-z0-9]+-->\n', '', tpl)
    tpl = re.sub(r'<!--@/?[a-z0-9]+-->', '', tpl)
    body = re.sub(r'\{\{([a-z0-9_]+)\}\}', lambda x: TX[x.group(1)], tpl)
    head = rd(BASE / 'head.html')
    return (head + '\n<!-- 대형 기획전 : 모양·기능은 공통 파일, 이 기획전 내용은 settings_bigevent.js -->\n'
            '<link rel="stylesheet" href="%s2026/00/bigevent/css/big_event.css?v=%s">\n\n' % (CDN, ENGINE_VER)
            + body + '\n'
            '<script src="%ssettings_bigevent.js?v=%s"></script>\n' % (meta['folder_url'], ver)
            + '<script src="%s2026/00/bigevent/js/big_event.js?v=%s"></script>\n' % (CDN, ENGINE_VER))

def cmd_make(folder, args):
    folder = Path(folder)
    S, TX, off, got, keep, meta = build(folder)
    ver = datetime.date.today().strftime('%y%m%d')
    out_dir = folder
    if (folder / 'settings_bigevent.js').exists() or (folder / 'index.html').exists():
        out_dir = folder / '_자동생성'
    s_js = settings_js(S, meta); idx = index_html(TX, off, meta, ver)
    wr(out_dir / 'settings_bigevent.js', s_js); wr(out_dir / 'index.html', idx)
    print('엑셀 :', Path(meta['xl']).name)
    print('만듦 : %s (settings_bigevent.js %d bytes · index.html %d bytes)' % (out_dir, len(s_js.encode()), len(idx.encode())))
    print('CDN 에 올릴 곳 : %s  ← settings_bigevent.js · img/…' % meta['folder_url'])
    print('\n[엑셀에서 읽음]'); [print('  ✔', g) for g in got]
    print('\n[빠진 섹션] ' + (', '.join(sorted(off)) if off else '없음'))
    print('\n[확인 필요 — 기본값(추석) 그대로]'); [print('  •', k) for k in dict.fromkeys(keep)]
    return out_dir

# ───────────────────────── check
def http(url, data=None, timeout=30):
    req = urllib.request.Request(url, data=data, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=timeout) as r: return r.status, r.read()

def cmd_check(folder, args):
    folder = Path(folder)
    S, TX, off, got, keep, meta = build(folder)
    bad = []
    pnos = set()
    for g in S['gift']: pnos.add(g['href'].split('=')[-1])
    if 'luckybag' not in off: pnos.add(S['luckybag']['href'].split('=')[-1])
    for key, days in S['dailyData'].items():
        for day, items in days.items():
            d = datetime.date.fromisoformat(day)
            if not (meta['start'] <= d <= meta['end']): bad.append('%s %s : 기간 밖 날짜' % (key, day))
            if len(items) != 12: bad.append('%s %s : 상품 %d개 (12개가 아님)' % (key, day, len(items)))
            if len(set(items)) != len(items): bad.append('%s %s : 같은 상품번호 겹침' % (key, day))
            pnos.update(items)
        miss = [(meta['start'] + datetime.timedelta(days=i)).isoformat() for i in range((meta['end'] - meta['start']).days + 1)]
        miss = [d for d in miss if d not in days]
        if days and miss: bad.append('%s : 날짜 빠짐 %s' % (key, ', '.join(miss)))
    print('상품번호 %d개 확인 중…' % len(pnos))
    dead = []
    for p in sorted(pnos):
        try:
            st, body = http('https://www.4xr.co.kr/shop/view.php?index_no=%s' % p)
            t = body.decode('utf-8', 'ignore')
            if 'og:title' not in t or '존재하지 않' in t or '판매가 종료' in t: dead.append(p)
        except Exception: dead.append(p)
    if dead: bad.append('상품 페이지가 없거나 판매 종료 : %s' % ', '.join(dead))
    # 이미지 : 폴더 img/ 에 있거나 CDN 에 있어야
    lost = []
    for rel, where in image_refs(S):
        u = S['img'] + rel
        if (folder / 'img' / rel).exists(): continue
        try:
            st, _ = http(u)
            if st == 200: continue
        except Exception: pass
        lost.append(rel)
    if lost: bad.append('이미지 없음(폴더 img/ 에도 CDN 에도) %d개 : %s' % (len(lost), ', '.join(lost[:12]) + (' …' if len(lost) > 12 else '')))
    print('\n[점검] ' + ('문제 없음 ✔' if not bad else '확인할 것 %d' % len(bad)))
    for b in bad: print('  ★', b)
    return not bad

# ───────────────────────── render
def find_playwright():
    for p in [Path.home() / 'Documents/4XR_기획전/.claude/skills/gihoekjeon-automation-system/scripts/node_modules/playwright', KIT / 'node_modules/playwright']:
        if p.exists(): return str(p)
    raise SystemExit('★ playwright 를 못 찾음')

def cmd_render(folder, args):
    folder = Path(folder)
    out_dir = folder / '_자동생성' if (folder / '_자동생성' / 'index.html').exists() else folder
    idx, st = out_dir / 'index.html', out_dir / 'settings_bigevent.js'
    if not idx.exists(): raise SystemExit('★ index.html 없음 — 먼저 make')
    shot = folder / '_점검'; shot.mkdir(exist_ok=True)
    url = args[0] if args else ''
    cmd = ['node', str(KIT / 'tools' / 'render.mjs'), find_playwright(), str(idx), str(st), str(KIT / 'engine'), str(folder / 'img'), str(shot), url]
    return subprocess.call(cmd) == 0

def main():
    if len(sys.argv) < 3 or sys.argv[1] not in ('init', 'make', 'check', 'render', 'all'):
        print(__doc__); sys.exit(1)
    cmd, folder, args = sys.argv[1], sys.argv[2], sys.argv[3:]
    if cmd == 'init': cmd_init(folder, args)
    elif cmd == 'make': cmd_make(folder, args)
    elif cmd == 'check': cmd_check(folder, args)
    elif cmd == 'render': cmd_render(folder, args)
    else:
        ok = cmd_check(folder, args); cmd_render(folder, args)

if __name__ == '__main__':
    main()
